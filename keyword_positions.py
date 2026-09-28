"""
Keyword position lookup for a sheet of (Page URL, Keyword, Country, Date Range).

Each row is looked up in this order and stops at the first source that has it:

  1. Google Search Console  impression-weighted average position for the exact
                            page + keyword + country over the row's date range.
  2. Ahrefs                 single-day organic rank for that exact URL + keyword.
  3. Semrush                top-100 rank for fortinet.com on the exact keyword,
                            matched to the exact page.
  4. Not ranking            none of the three. Notes then says which other page
                            of the site ranks instead (if any) and what the page
                            itself ranks for, so the row is still useful.

The Source column says where every number came from. Only GSC numbers are
weekly averages; Ahrefs and Semrush are point-in-time ranks, so never average
them together.

Setup (once):
    pip install pandas openpyxl requests google-auth google-api-python-client

Keys: put them in a file called keys.env next to this script:
    SERVICE_ACCOUNT_FILE=fortinet-gsc-api-c8ebcebdce98.json
    SITE_URL=https://www.fortinet.com
    AHREFS_API_TOKEN=xxxx
    SEMRUSH_API_KEY=xxxx
(or set them as environment variables). A missing key just skips that source.

Run:
    python keyword_positions.py --input Keyword_input.csv
    python keyword_positions.py --input Keyword_input.csv --last-week
    python keyword_positions.py --input Keyword_input.csv --output Fortinet_positions.xlsx
    python keyword_positions.py --input Keyword_input.csv --last-week --master Fortinet_Keyword_Positions.xlsx
    python keyword_positions.py --input Keyword_input.csv --week-start 2026-09-13 --master Fortinet_Keyword_Positions.xlsx
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urlsplit, urlunsplit

import pandas as pd
import requests

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-7s %(message)s",
                    datefmt="%H:%M:%S")
log = logging.getLogger("positions")


# --------------------------------------------------------------------------
# Config
# --------------------------------------------------------------------------

def load_keys_file(path: Path) -> None:
    """Read KEY=value lines from keys.env into the environment (env vars win)."""
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


HERE = Path(__file__).resolve().parent
load_keys_file(HERE / "keys.env")

SERVICE_ACCOUNT_FILE = os.environ.get("SERVICE_ACCOUNT_FILE", "fortinet-gsc-api-c8ebcebdce98.json")
SITE_URL = os.environ.get("SITE_URL", "https://www.fortinet.com")
AHREFS_API_TOKEN = os.environ.get("AHREFS_API_TOKEN", "")
SEMRUSH_API_KEY = os.environ.get("SEMRUSH_API_KEY", "")

AHREFS_ENDPOINT = "https://api.ahrefs.com/v3/site-explorer/organic-keywords"
SEMRUSH_ENDPOINT = "https://api.semrush.com/"          # v3 keys
SEMRUSH_MCP_ENDPOINT = "https://mcp.semrush.com/v2/mcp"  # v4 keys (semrtkn-...)
GSC_SCOPES = ["https://www.googleapis.com/auth/webmasters.readonly"]
SAFE_LAG_DAYS = 3  # GSC data is final ~3 days back

# name/code -> (GSC alpha-3, 2-letter code used by Ahrefs/Semrush)
COUNTRY_MAP: Dict[str, Tuple[str, str]] = {
    "united states": ("usa", "us"), "us": ("usa", "us"), "usa": ("usa", "us"),
    "united kingdom": ("gbr", "gb"), "uk": ("gbr", "gb"), "gb": ("gbr", "gb"),
    "india": ("ind", "in"), "in": ("ind", "in"),
    "canada": ("can", "ca"), "ca": ("can", "ca"),
    "australia": ("aus", "au"), "au": ("aus", "au"),
    "germany": ("deu", "de"), "de": ("deu", "de"),
    "france": ("fra", "fr"), "fr": ("fra", "fr"),
    "japan": ("jpn", "jp"), "jp": ("jpn", "jp"),
    "singapore": ("sgp", "sg"), "sg": ("sgp", "sg"),
    "brazil": ("bra", "br"), "br": ("bra", "br"),
    "mexico": ("mex", "mx"), "mx": ("mex", "mx"),
    "spain": ("esp", "es"), "es": ("esp", "es"),
    "italy": ("ita", "it"), "it": ("ita", "it"),
    "netherlands": ("nld", "nl"), "nl": ("nld", "nl"),
    "united arab emirates": ("are", "ae"), "ae": ("are", "ae"),
}
SEMRUSH_DB_FIX = {"gb": "uk"}  # Semrush calls the UK database "uk"

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

COLUMN_ALIASES = {
    "url": {"page urls", "page url", "url", "urls", "page", "landing page"},
    "keyword": {"keywords", "keyword", "query", "queries", "search term"},
    "country": {"country", "countries", "market", "location"},
    "date_range": {"date range", "daterange", "date", "dates", "period", "week"},
}


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def normalize_url(raw: str) -> str:
    s = str(raw).strip()
    if not s:
        return ""
    if not s.startswith(("http://", "https://")):
        s = "https://" + s
    p = urlsplit(s)
    return urlunsplit((p.scheme, p.netloc, p.path.rstrip("/") or "/", "", ""))


def url_key(url: str) -> str:
    """Compare URLs ignoring http/https, www, case and trailing slash."""
    p = urlsplit(normalize_url(url))
    host = p.netloc.lower()
    host = host[4:] if host.startswith("www.") else host
    return host + (p.path.rstrip("/") or "/").lower()


def host_of(url: str) -> str:
    h = urlsplit(normalize_url(url)).netloc.lower()
    return h[4:] if h.startswith("www.") else h


def url_variants(url: str) -> List[str]:
    base = normalize_url(url)
    p = urlsplit(base)
    alt_host = p.netloc[4:] if p.netloc.startswith("www.") else "www." + p.netloc
    alt = urlunsplit((p.scheme, alt_host, p.path, "", ""))
    out = [base, base + "/", alt, alt + "/"] if p.path != "/" else [base, alt]
    return list(dict.fromkeys(out))


def resolve_country(raw: Any) -> Tuple[Optional[str], Optional[str], str]:
    s = str(raw if raw is not None else "").strip()
    if not s or s.lower() in {"nan", "all", "worldwide", "global"}:
        return None, None, "Worldwide"
    for key in (s.lower(),
                (re.search(r"\(([A-Za-z]{2,3})\)", s) or [None, ""])[1].lower(),
                re.sub(r"\s*\([^)]*\)", "", s).strip().lower()):
        if key and key in COUNTRY_MAP:
            a3, a2 = COUNTRY_MAP[key]
            return a3, a2, s
    log.warning("Unknown country %r, using worldwide. Add it to COUNTRY_MAP.", s)
    return None, None, s


def format_range(a: date, b: date) -> str:
    if a.month == b.month and a.year == b.year:
        return f"{a:%b} {a.day}-{b.day}, {b.year}"
    return f"{a:%b} {a.day} - {b:%b} {b.day}, {b.year}"


def last_week_range(today: Optional[date] = None) -> Tuple[date, date]:
    """Most recent complete Sun-Sat week that is past the GSC lag."""
    cutoff = (today or date.today()) - timedelta(days=SAFE_LAG_DAYS)
    end = cutoff - timedelta(days=(cutoff.weekday() + 2) % 7)
    return end - timedelta(days=6), end


def parse_date_range(raw: Any, fallback: Tuple[date, date]) -> Tuple[date, date, str]:
    s = str(raw if raw is not None else "").strip()
    if not s or s.lower() in {"nan", "last week", "lastweek", "last_week"}:
        return fallback[0], fallback[1], format_range(*fallback)
    m = re.search(r"(\d{4}-\d{2}-\d{2})\s*(?:to|-|–|—)\s*(\d{4}-\d{2}-\d{2})", s)
    if m:
        return (datetime.strptime(m.group(1), "%Y-%m-%d").date(),
                datetime.strptime(m.group(2), "%Y-%m-%d").date(), s)
    m = re.search(r"([A-Za-z]{3,9})\.?\s*(\d{1,2})\s*(?:-|–|—|to)\s*([A-Za-z]{3,9})\.?\s*(\d{1,2})\s*,?\s*(\d{4})", s)
    if m and MONTHS.get(m.group(1)[:3].lower()) and MONTHS.get(m.group(3)[:3].lower()):
        y = int(m.group(5))
        a = date(y, MONTHS[m.group(1)[:3].lower()], int(m.group(2)))
        b = date(y, MONTHS[m.group(3)[:3].lower()], int(m.group(4)))
        if b < a:
            b = date(y + 1, b.month, b.day)
        return a, b, s
    m = re.search(r"([A-Za-z]{3,9})\.?\s*(\d{1,2})\s*(?:-|–|—|to)\s*(\d{1,2})\s*,?\s*(\d{4})", s)
    if m and MONTHS.get(m.group(1)[:3].lower()):
        mo, y = MONTHS[m.group(1)[:3].lower()], int(m.group(4))
        return date(y, mo, int(m.group(2))), date(y, mo, int(m.group(3))), s
    log.warning("Could not read date range %r, using last week.", s)
    return fallback[0], fallback[1], s


def find_columns(df: pd.DataFrame) -> Dict[str, str]:
    lookup = {str(c).strip().lower().lstrip("﻿"): c for c in df.columns}
    found = {k: lookup[a] for k, aliases in COLUMN_ALIASES.items()
             for a in aliases if a in lookup}
    missing = [c for c in ("url", "keyword") if c not in found]
    if missing:
        raise SystemExit(f"Input is missing column(s) {missing}. Found: {list(df.columns)}")
    return found


# --------------------------------------------------------------------------
# 1. Google Search Console
# --------------------------------------------------------------------------

class GSC:
    def __init__(self, key_file: str, site: str):
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        creds = service_account.Credentials.from_service_account_file(key_file, scopes=GSC_SCOPES)
        self.service = build("searchconsole", "v1", credentials=creds, cache_discovery=False)
        self.site = site

    def lookup(self, url: str, keyword: str, country: Optional[str],
               start: date, end: date) -> Optional[Dict[str, Any]]:
        kw = keyword.strip().lower()  # GSC stores queries lowercased
        for candidate in url_variants(url):
            filters = [{"dimension": "page", "operator": "equals", "expression": candidate},
                       {"dimension": "query", "operator": "equals", "expression": kw}]
            if country:
                filters.append({"dimension": "country", "operator": "equals", "expression": country})
            body = {"startDate": start.isoformat(), "endDate": end.isoformat(),
                    "dimensions": ["page", "query"], "rowLimit": 5, "dataState": "final",
                    "dimensionFilterGroups": [{"filters": filters}]}
            try:
                rows = self.service.searchanalytics().query(siteUrl=self.site, body=body).execute().get("rows", [])
            except Exception as exc:  # noqa: BLE001
                log.error("GSC failed for %r: %s", keyword, exc)
                return None
            if rows:
                r = rows[0]
                return {"position": round(float(r["position"]), 2),
                        "impressions": int(r.get("impressions", 0)), "clicks": int(r.get("clicks", 0))}
        return None


# --------------------------------------------------------------------------
# 2. Ahrefs
# --------------------------------------------------------------------------

class Ahrefs:
    def __init__(self, token: str):
        self.token = token
        self.session = requests.Session()
        self.session.headers.update({"Authorization": f"Bearer {token}", "Accept": "application/json"})

    def lookup(self, url: str, keyword: str, country: Optional[str], on: date) -> Optional[Dict[str, Any]]:
        if not self.token:
            return None
        kw = keyword.strip().lower().replace('"', '\\"')
        params = {"target": normalize_url(url), "mode": "exact", "date": on.isoformat(),
                  "select": "keyword,best_position,best_position_url,volume",
                  "where": f'{{"field":"keyword","is":["eq","{kw}"]}}', "limit": 5}
        if country:
            params["country"] = country
        try:
            resp = self.session.get(AHREFS_ENDPOINT, params=params, timeout=45)
            if resp.status_code != 200:
                log.error("Ahrefs HTTP %s for %r: %s", resp.status_code, keyword, resp.text[:150])
                return None
            rows = resp.json().get("keywords", [])
        except Exception as exc:  # noqa: BLE001
            log.error("Ahrefs failed for %r: %s", keyword, exc)
            return None
        if rows and rows[0].get("best_position") is not None:
            return {"position": float(rows[0]["best_position"]), "volume": rows[0].get("volume")}
        return None


# --------------------------------------------------------------------------
# 3. Semrush
# --------------------------------------------------------------------------

class Semrush:
    """domain_organic for the exact keyword; url_organic to explain misses.
    Costs ~10 API units per returned row."""

    def __init__(self, key: str):
        self.key = key
        self.session = requests.Session()
        self.units = 0
        self._cache: Dict[Tuple[str, str, str], Optional[List[Dict[str, str]]]] = {}

    # ---- transport -----------------------------------------------------
    @property
    def uses_mcp(self) -> bool:
        """v4 keys (semrtkn-...) only work through Semrush's MCP server."""
        return self.key.startswith("semrtkn")

    def _get(self, params: Dict[str, Any]) -> Optional[List[Dict[str, str]]]:
        """None = request failed, [] = nothing found. Rows are dicts keyed by the
        CSV header Semrush returns (Keyword, Position, Url, Timestamp, Search Volume)."""
        text = self._mcp(params) if self.uses_mcp else self._v3(params)
        if text is None:
            return None
        if text == "":
            return []
        lines = [l for l in text.strip().splitlines() if l.strip()]
        if len(lines) < 2:
            return []
        head = [h.strip() for h in lines[0].split(";")]
        rows = [dict(zip(head, [c.strip() for c in l.split(";")])) for l in lines[1:]]
        self.units += 10 * len(rows)
        return rows

    def _v3(self, params: Dict[str, Any]) -> Optional[str]:
        params = dict(params, key=self.key)
        for attempt in range(3):
            try:
                resp = self.session.get(SEMRUSH_ENDPOINT, params=params, timeout=60)
            except Exception as exc:  # noqa: BLE001
                log.warning("Semrush error (attempt %d): %s", attempt + 1, exc)
                time.sleep(2 * (attempt + 1))
                continue
            if resp.status_code == 429 or resp.status_code >= 500:
                time.sleep(3 * (attempt + 1))
                continue
            text = resp.text.strip()
            if text.startswith("ERROR 50"):
                return ""
            if text.startswith("ERROR") or text.startswith("{"):
                log.error("Semrush: %s", text[:200])
                return None
            return text
        return None

    # ---- MCP (v4 key) ---------------------------------------------------
    def _mcp_post(self, payload: Dict[str, Any]) -> Tuple[Optional[Dict[str, Any]], Any]:
        headers = {"Authorization": f"Apikey {self.key}", "Content-Type": "application/json",
                   "Accept": "application/json, text/event-stream",
                   "MCP-Protocol-Version": "2025-06-18"}
        if getattr(self, "_sid", None):
            headers["Mcp-Session-Id"] = self._sid
        resp = self.session.post(SEMRUSH_MCP_ENDPOINT, json=payload, headers=headers, timeout=90)
        if resp.status_code in (401, 403):
            raise PermissionError(f"Semrush MCP rejected the key (HTTP {resp.status_code}): {resp.text[:200]}")
        resp.raise_for_status()
        if resp.headers.get("Mcp-Session-Id"):
            self._sid = resp.headers["Mcp-Session-Id"]
        if "id" not in payload or not resp.content:
            return None, resp
        body = resp.text
        if "text/event-stream" in resp.headers.get("Content-Type", ""):
            import json
            msgs = [json.loads(l[5:].strip()) for l in body.splitlines() if l.startswith("data:") and l[5:].strip()]
            msg = next((m for m in msgs if m.get("id") == payload["id"]), msgs[-1] if msgs else {})
        else:
            msg = resp.json()
        return msg, resp

    def _mcp_init(self) -> None:
        if getattr(self, "_ready", False):
            return
        self._sid, self._rid = None, 0
        msg, _ = self._mcp_post({"jsonrpc": "2.0", "id": 0, "method": "initialize", "params": {
            "protocolVersion": "2025-06-18", "capabilities": {},
            "clientInfo": {"name": "leadwalnut-keyword-positions", "version": "1.0"}}})
        if not msg or "error" in msg:
            raise RuntimeError(f"Semrush MCP initialize failed: {msg}")
        self._mcp_post({"jsonrpc": "2.0", "method": "notifications/initialized"})
        self._ready = True

    def _mcp(self, params: Dict[str, Any]) -> Optional[str]:
        """Translate the v3-style params into the MCP 'resource_organic' report."""
        import json
        if params["type"] == "domain_organic":
            kw = params["display_filter"].split("|", 3)[3]
            report = {"target": params["domain"], "database": params["database"],
                      "display_filter": [{"field": "keyword", "operation": "equals", "sign": "+", "value": kw}],
                      "export_columns": ["keyword", "position", "url", "timestamp"],
                      "display_limit": params.get("display_limit", 20)}
        else:  # url_organic
            report = {"target": params["url"], "database": params["database"],
                      "display_sort": "traffic_desc", "export_columns": ["keyword", "position", "volume"],
                      "display_limit": params.get("display_limit", 3)}
        for attempt in range(3):
            try:
                self._mcp_init()
                self._rid += 1
                msg, _ = self._mcp_post({"jsonrpc": "2.0", "id": self._rid, "method": "tools/call",
                                         "params": {"name": "execute_report",
                                                    "arguments": {"report": "resource_organic", "params": report}}})
            except PermissionError as exc:
                log.error("%s", exc)
                return None
            except Exception as exc:  # noqa: BLE001
                log.warning("Semrush MCP error (attempt %d): %s", attempt + 1, exc)
                self._ready = False
                time.sleep(3 * (attempt + 1))
                continue
            if msg and "error" in msg and "NOTHING FOUND" in str(msg["error"]):
                return ""  # Semrush's normal "no rankings" answer, not a failure
            if not msg or "error" in msg:
                log.error("Semrush MCP error: %s", (msg or {}).get("error"))
                return None
            result = msg.get("result", {})
            text = "".join(c.get("text", "") for c in result.get("content", []) if c.get("type") == "text")
            if "NOTHING FOUND" in text:
                return ""
            if result.get("isError"):
                log.error("Semrush MCP: %s", text[:200])
                return None
            try:
                return json.loads(text).get("data", "")
            except ValueError:
                return text
        return None

    @staticmethod
    def db(country: Optional[str]) -> str:
        c = (country or "us").lower()
        return SEMRUSH_DB_FIX.get(c, c)

    def lookup(self, url: str, keyword: str, country: Optional[str]) -> Optional[Dict[str, Any]]:
        if not self.key:
            return None
        kw, dom, db = keyword.strip().lower(), host_of(url), self.db(country)
        if (kw, dom, db) not in self._cache:
            self._cache[(kw, dom, db)] = self._get({
                "type": "domain_organic", "domain": dom, "database": db,
                "display_filter": f"+|Ph|Eq|{kw}", "export_columns": "Ph,Po,Ur,Ts",
                "display_limit": 20})
        rows = self._cache[(kw, dom, db)]
        if rows is None:
            return {"status": "error"}
        target = url_key(url)
        rows = sorted(rows, key=lambda r: int(r.get("Position") or 999))
        mine = [r for r in rows if url_key(r.get("Url", "")) == target]
        other = next((r for r in rows if url_key(r.get("Url", "")) != target), None)
        when = ""
        ts = (rows[0].get("Timestamp") if rows else "") or ""
        if ts.isdigit():
            when = datetime.fromtimestamp(int(ts), timezone.utc).strftime("%Y-%m-%d")
        return {"status": "found" if mine else "not_found",
                "position": float(mine[0]["Position"]) if mine else None,
                "date": when, "other": other}

    def page_keywords(self, url: str, country: Optional[str], n: int = 3) -> List[Dict[str, str]]:
        if not self.key:
            return []
        return self._get({"type": "url_organic", "url": normalize_url(url),
                          "database": self.db(country), "display_sort": "tr_desc",
                          "export_columns": "Ph,Po,Nq", "display_limit": n}) or []


# --------------------------------------------------------------------------
# Master workbook: one tab per week + Summary + Trend
# --------------------------------------------------------------------------

POSITION_COLS = ["Page URL", "Keyword", "Country", "Date Range", "Average Position", "Source",
                 "Data Date", "Impressions", "Clicks", "Notes"]
FILLS = {"GSC": "E2EFDA", "Ahrefs": "FFF2CC", "Semrush": "DDEBF7", "Not ranking": "FCE4E4"}


def week_tab_name(start: date, end: date) -> str:
    """'Sep 13-19' or 'Sep 27-Oct 3' (Excel tab names max 31 chars, no / : etc.)."""
    if start.month == end.month:
        return f"{start:%b} {start.day}-{end.day}"
    return f"{start:%b} {start.day}-{end:%b} {end.day}"


def _week_start_of(df: pd.DataFrame) -> date:
    first = str(df["Date Range"].dropna().iloc[0]) if "Date Range" in df and df["Date Range"].notna().any() else ""
    return parse_date_range(first, (date.min, date.min))[0]


def update_master(week_df: pd.DataFrame, start: date, end: date, path: str) -> str:
    """Add (or replace) this week's tab in the master workbook and rebuild Summary + Trend."""
    weeks: Dict[str, pd.DataFrame] = {}
    if Path(path).exists():
        for name, sheet in pd.read_excel(path, sheet_name=None).items():
            if name not in ("Summary", "Trend") and "Keyword" in sheet.columns:
                weeks[name] = sheet
    tab = week_tab_name(start, end)
    weeks[tab] = week_df
    order = sorted(weeks, key=lambda n: _week_start_of(weeks[n]))  # oldest -> newest

    summary_rows = []
    for name in order:
        w = weeks[name]
        src = w["Source"].astype(str)
        gsc_pos = pd.to_numeric(w.loc[src == "GSC", "Average Position"], errors="coerce")
        allpos = pd.to_numeric(w["Average Position"], errors="coerce")
        summary_rows.append({
            "Week": name, "Keywords": len(w),
            "GSC": int((src == "GSC").sum()), "Ahrefs": int((src == "Ahrefs").sum()),
            "Semrush": int((src == "Semrush").sum()),
            "Other sources": int((~src.isin(["GSC", "Ahrefs", "Semrush", "Not ranking"])).sum()),
            "Not ranking": int((src == "Not ranking").sum()),
            "Avg position (GSC rows)": round(gsc_pos.mean(), 2) if gsc_pos.notna().any() else None,
            "In top 3": int((allpos <= 3).sum()), "In top 10": int((allpos <= 10).sum()),
        })
    summary = pd.DataFrame(summary_rows)

    # Trend: one row per URL+keyword, one column per week (oldest -> newest)
    trend = None
    for name in order:
        w = weeks[name][["Page URL", "Keyword", "Average Position"]].copy()
        w["Keyword"] = w["Keyword"].astype(str)
        w = w.drop_duplicates(["Page URL", "Keyword"]).rename(columns={"Average Position": name})
        trend = w if trend is None else trend.merge(w, on=["Page URL", "Keyword"], how="outer")
    if len(order) >= 2:
        prev, last = pd.to_numeric(trend[order[-2]], errors="coerce"), pd.to_numeric(trend[order[-1]], errors="coerce")
        trend["Change vs last week"] = (prev - last).round(2)  # positive = moved up

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        trend.to_excel(xw, sheet_name="Trend", index=False)
        for name in reversed(order):  # newest week first after Summary/Trend
            weeks[name].reindex(columns=[c for c in POSITION_COLS if c in weeks[name].columns] +
                                [c for c in weeks[name].columns if c not in POSITION_COLS]) \
                .to_excel(xw, sheet_name=name, index=False)
        _style_book(xw.book)
    log.info("Master workbook %s updated: tab '%s' (%d weeks total)", path, tab, len(order))
    return tab


def _style_book(book) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for ws in book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.alignment = Alignment(vertical="center", wrap_text=False)
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)
        ws.freeze_panes = "C2" if ws.title == "Trend" else "A2"
        headers = [c.value for c in ws[1]]
        if "Source" in headers:
            sc = headers.index("Source") + 1
            for r in range(2, ws.max_row + 1):
                color = FILLS.get(ws.cell(r, sc).value)
                if color:
                    for c in range(1, ws.max_column + 1):
                        ws.cell(r, c).fill = PatternFill("solid", fgColor=color)
        if "Change vs last week" in headers:
            cc = headers.index("Change vs last week") + 1
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, cc).value
                if isinstance(v, (int, float)) and v:
                    ws.cell(r, cc).font = Font(bold=True, color="1E7B34" if v > 0 else "C00000")


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def run(args: argparse.Namespace) -> int:
    src = Path(args.input)
    if not src.exists():
        raise SystemExit(f"Input file not found: {src}")
    df = (pd.read_csv(src, sep=None, engine="python", encoding="utf-8-sig")
          if src.suffix.lower() in {".csv", ".tsv", ".txt"} else pd.read_excel(src))
    if df.empty:
        raise SystemExit("Input file has no rows.")
    cols = find_columns(df)
    week = last_week_range()
    if args.week_start:
        ws_ = datetime.strptime(args.week_start, "%Y-%m-%d").date()
        week = (ws_, ws_ + timedelta(days=6))
        args.last_week = True  # use this week for every row
    log.info("Loaded %d rows from %s", len(df), src.name)

    gsc = None
    try:
        gsc = GSC(SERVICE_ACCOUNT_FILE, SITE_URL)
        log.info("GSC ready for %s", SITE_URL)
    except Exception as exc:  # noqa: BLE001
        log.error("GSC not available (%s). Continuing with Ahrefs and Semrush.", exc)
    ahrefs = Ahrefs(AHREFS_API_TOKEN)
    semrush = Semrush(SEMRUSH_API_KEY)
    for name, ok in (("Ahrefs", AHREFS_API_TOKEN), ("Semrush", SEMRUSH_API_KEY)):
        log.info("%s %s", name, "ready" if ok else "key not set, skipped")

    counts = {"GSC": 0, "Ahrefs": 0, "Semrush": 0, "Not ranking": 0}
    out: List[Dict[str, Any]] = []

    for n, (_, row) in enumerate(df.iterrows(), start=1):
        url = str(row[cols["url"]]).strip()
        keyword = str(row[cols["keyword"]]).strip()
        if not url or url.lower() == "nan" or not keyword or keyword.lower() == "nan":
            continue
        a3, a2, country_label = resolve_country(row[cols["country"]] if "country" in cols else "")
        if args.last_week:
            start, end = week
            label = format_range(start, end)
        else:
            start, end, label = parse_date_range(row[cols["date_range"]] if "date_range" in cols else "", week)

        rec = {"Page URL": url, "Keyword": keyword, "Country": country_label, "Date Range": label,
               "Average Position": None, "Source": "Not ranking", "Data Date": "",
               "Impressions": None, "Clicks": None, "Notes": ""}

        hit = gsc.lookup(url, keyword, a3, start, end) if gsc else None
        if hit:
            rec.update({"Average Position": hit["position"], "Source": "GSC",
                        "Data Date": label, "Impressions": hit["impressions"], "Clicks": hit["clicks"]})
        else:
            alt = ahrefs.lookup(url, keyword, a2, end)
            if alt:
                rec.update({"Average Position": alt["position"], "Source": "Ahrefs",
                            "Data Date": end.isoformat(),
                            "Notes": "Ahrefs single-day rank" + (f"; volume {alt['volume']}" if alt.get("volume") else "")})
            else:
                sem = semrush.lookup(url, keyword, a2)
                notes = ["No GSC impressions", "no Ahrefs rank"]
                if sem and sem["status"] == "found":
                    rec.update({"Average Position": sem["position"], "Source": "Semrush",
                                "Data Date": sem["date"]})
                    notes = ["Semrush top-100 rank"]
                    if sem.get("other"):
                        notes.append(f"another page also ranks #{sem['other']['Position']}: {sem['other']['Url']}")
                elif sem and sem["status"] == "error":
                    notes.append("Semrush request failed")
                elif sem:
                    if sem.get("other"):
                        notes.append(f"a different page ranks #{sem['other']['Position']}: {sem['other']['Url']}")
                    else:
                        notes.append("not in Semrush top 100")
                    top = semrush.page_keywords(url, a2)
                    if top:
                        notes.append("this page ranks for " + ", ".join(
                            f"'{t['Keyword']}' #{t['Position']} (vol {t['Search Volume']})" for t in top))
                    else:
                        notes.append("Semrush has no rankings for this page at all, check indexing in GSC")
                else:
                    notes.append("Semrush key not set")
                rec["Notes"] = "; ".join(notes)

        counts[rec["Source"]] += 1
        out.append(rec)
        log.info("[%d/%d] %-40s -> %s (%s)", n, len(df), keyword[:40],
                 rec["Average Position"], rec["Source"])
        time.sleep(args.delay)

    result = pd.DataFrame(out)
    output = args.output or f"{src.stem}_positions_{date.today():%Y%m%d}.xlsx"
    write_workbook(result, counts, output)
    log.info("-" * 60)
    log.info("Wrote %d rows to %s", len(result), output)
    if args.master:
        if not args.last_week:
            log.warning("--master works best with --last-week or --week-start (one week per tab)")
        update_master(result, week[0], week[1], args.master)
    log.info("GSC: %d   Ahrefs: %d   Semrush: %d   Not ranking: %d",
             counts["GSC"], counts["Ahrefs"], counts["Semrush"], counts["Not ranking"])
    if SEMRUSH_API_KEY:
        log.info("Semrush API units used (approx): %d", semrush.units)
    return 0


def write_workbook(df: pd.DataFrame, counts: Dict[str, int], path: str) -> None:
    """Positions sheet (colour-coded by source) + a Summary sheet."""
    from openpyxl.styles import Alignment, Font, PatternFill
    fills = {"GSC": "E2EFDA", "Ahrefs": "FFF2CC", "Semrush": "DDEBF7", "Not ranking": "FCE4E4"}
    summary = pd.DataFrame([
        {"Source": "GSC", "Rows": counts["GSC"], "What the number means": "Impression-weighted average over the date range"},
        {"Source": "Ahrefs", "Rows": counts["Ahrefs"], "What the number means": "Single-day rank at the end of the date range"},
        {"Source": "Semrush", "Rows": counts["Semrush"], "What the number means": "Top-100 rank from Semrush's latest crawl (see Data Date)"},
        {"Source": "Not ranking", "Rows": counts["Not ranking"], "What the number means": "Not in any source's top 100; see Notes"},
        {"Source": "Total", "Rows": sum(counts.values()), "What the number means": ""},
    ])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Positions", index=False)
        summary.to_excel(xw, sheet_name="Summary", index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3864")
                cell.alignment = Alignment(vertical="center")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)
            ws.freeze_panes = "A2"
        ws = xw.book["Positions"]
        src_col = list(df.columns).index("Source") + 1
        for r in range(2, ws.max_row + 1):
            color = fills.get(ws.cell(r, src_col).value)
            if color:
                for c in range(1, ws.max_column + 1):
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=color)


def semrush_selftest() -> int:
    """Quick check that the Semrush key works: looks up 'proxy server' for fortinet.com."""
    if not SEMRUSH_API_KEY:
        raise SystemExit("SEMRUSH_API_KEY is not set (keys.env or environment).")
    sem = Semrush(SEMRUSH_API_KEY)
    log.info("Key type: %s", "v4 via Semrush MCP" if sem.uses_mcp else "v3 via api.semrush.com")
    res = sem.lookup("https://www.fortinet.com/resources/cyberglossary/proxy-server", "proxy server", "us")
    log.info("Result: %s", res)
    ok = bool(res) and res.get("status") in ("found", "not_found")
    log.info("Semrush is %s", "WORKING" if ok else "NOT working, see errors above")
    return 0 if ok else 1


def main() -> int:
    if "--semrush-test" in sys.argv:
        return semrush_selftest()
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--input", "-i", required=True, help="Keyword sheet (.csv or .xlsx)")
    p.add_argument("--output", "-o", help="Output .xlsx (default: <input>_positions_<today>.xlsx)")
    p.add_argument("--last-week", action="store_true",
                   help="Ignore the Date Range column and use the last complete Sun-Sat week")
    p.add_argument("--delay", type=float, default=0.2, help="Seconds between rows")
    p.add_argument("--week-start", metavar="YYYY-MM-DD",
                   help="Run a specific Sun-Sat week, e.g. 2026-09-13 (for backfilling)")
    p.add_argument("--master", metavar="XLSX",
                   help="Also add this week as a tab in a running master workbook")
    return run(p.parse_args())


if __name__ == "__main__":
    sys.exit(main())
